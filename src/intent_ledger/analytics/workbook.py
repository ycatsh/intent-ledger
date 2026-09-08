import csv
import io
import re
import zipfile
from calendar import month_name, monthrange
from collections import defaultdict
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from intent_ledger.accounting.accounts import get_account, get_ledger_entries
from intent_ledger.accounting.projects import (
    get_project,
    get_project_category_summary,
    get_project_inflow_summary,
    get_project_payee_summary,
    get_project_summary,
    get_project_transactions,
)
from intent_ledger.analytics.reports import (
    get_category_summary,
    get_inflow_summary,
    get_monthly_summary,
    get_payee_summary,
)
from intent_ledger.config import DATA_DIR
from intent_ledger.db import db
from intent_ledger.domain.money import Money
from intent_ledger.settings import get_export_format, today

FINANCE_EXPORT_DIR = DATA_DIR / "finance" / "exports"

HEADER_FILL = PatternFill(
    fill_type="solid",
    start_color="217346",
    end_color="217346",
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
)

SPLIT_FILL = PatternFill(
    fill_type="solid",
    start_color="FBF3DB",
    end_color="FBF3DB",
)

CATEGORY_FILL = PatternFill(
    fill_type="solid",
    start_color="F2F4F6",
    end_color="F2F4F6",
)

ACCOUNTING_FORMAT = "+#,##0.00;-#,##0.00"

SUBTOTAL_BORDER = Border(top=Side(style="thin", color="C9CFD6"))
GRAND_TOTAL_BORDER = Border(top=Side(style="thin", color="C9CFD6"))

TRANSACTIONS_HEADERS = ["Date", "Payee", "Category", "Type", "Debit", "Credit", "Description"]


def _export_path(filename):
    FINANCE_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    return FINANCE_EXPORT_DIR / filename


def _save_workbook(wb, stem):
    """Save `wb` under `stem` in the app's configured export format
    (Settings > Export format). CSV can't hold more than one sheet per file,
    so a multi-sheet report becomes a .zip of one .csv per sheet instead -
    still a single download, same data, no formatting.
    """
    with db.transaction() as conn:
        fmt = get_export_format(conn)

    if fmt == "csv":
        return _save_as_csv(wb, stem)

    path = _export_path(f"{stem}.xlsx")
    wb.save(path)
    return path


def _save_as_csv(wb, stem):
    path = _export_path(f"{stem}.zip")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for sheet in wb.worksheets:
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
            archive.writestr(f"{_slugify(sheet.title)}.csv", buffer.getvalue())

    return path


def _collapse_column(sheet, index):
    letter = get_column_letter(index)
    sheet.column_dimensions[letter].outlineLevel = 1
    sheet.column_dimensions[letter].hidden = True


def _round_money(value):
    return round(value, 2) or 0.0


def _positive_amounts(rows):
    return [{**r, "amount": -r["amount"]} for r in rows]


def export_monthly_report(year, month):
    stem = f"exp-report_{year}_{month:02d}"

    start = f"{year}-{month:02d}-01"
    end = f"{year}-{month:02d}-{monthrange(year, month)[1]:02d}"

    summary = get_monthly_summary(year, month)
    categories = get_category_summary(year, month)
    payees = get_payee_summary(year, month)
    inflows = get_inflow_summary(year, month)
    transactions = _transactions(start, end)

    wb = Workbook()

    _summary_sheet(
        wb.active,
        f"{month_name[month]} {year}",
        summary,
        categories,
    )

    txn_sheet = wb.create_sheet("Transactions")
    _transactions_sheet(txn_sheet, TRANSACTIONS_HEADERS, transactions)
    _collapse_column(txn_sheet, len(TRANSACTIONS_HEADERS))

    _table_sheet(
        wb.create_sheet("Payees"),
        ["Payee", "Amount", "Transactions", "%"],
        [
            [
                x["payee"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in payees
        ],
    )

    _table_sheet(
        wb.create_sheet("Inflows"),
        ["Source", "Amount", "Transactions", "%"],
        [
            [
                x["source"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in inflows
        ],
    )

    _style(wb)

    return _save_workbook(wb, stem)


def export_yearly_report(year):
    stem = f"exp-report_{year}"

    wb = Workbook()

    summary = get_monthly_summary(year, month=None)
    categories = get_category_summary(year, month=None)
    payees = get_payee_summary(year, month=None)
    inflows = get_inflow_summary(year, month=None)

    _summary_sheet(
        wb.active,
        str(year),
        summary,
        categories,
    )

    txn_sheet = wb.create_sheet("Transactions")
    _transactions_sheet(
        txn_sheet,
        TRANSACTIONS_HEADERS,
        _transactions(
            f"{year}-01-01",
            f"{year}-12-31",
        ),
    )
    _collapse_column(txn_sheet, len(TRANSACTIONS_HEADERS))

    _table_sheet(
        wb.create_sheet("Payees"),
        ["Payee", "Amount", "Transactions", "%"],
        [
            [
                x["payee"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in payees
        ],
    )

    _table_sheet(
        wb.create_sheet("Inflows"),
        ["Source", "Amount", "Transactions", "%"],
        [
            [
                x["source"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in inflows
        ],
    )

    _style(wb)

    return _save_workbook(wb, stem)


def export_account_statement(account_id, search=None):
    account = get_account(account_id)
    if account is None:
        raise ValueError(f"Account {account_id} not found.")

    full = get_ledger_entries(account_id, sort="date_asc")

    running = 0.0
    balances = {}
    for t in full:
        running = _round_money(running + t["amount"])
        balances[t["transaction_hash"]] = running

    filtered = [
        t
        for t in get_ledger_entries(account_id, search=search, sort="date_asc")
        if t["category_type"] != "equity"
    ]

    if filtered:
        first_index = next(
            i for i, t in enumerate(full) if t["transaction_hash"] == filtered[0]["transaction_hash"]
        )
        opening_balance = balances[full[first_index - 1]["transaction_hash"]] if first_index > 0 else 0.0
        closing_balance = balances[filtered[-1]["transaction_hash"]]
        date_range = f"{filtered[0]['date'][:10]}_to_{filtered[-1]['date'][:10]}"
    else:
        opening_balance = closing_balance = 0.0
        date_range = "no-results"

    stem = f"statement_{_slugify(account['name'])}_{date_range}"

    total_debit = _round_money(-sum(t["amount"] for t in filtered if t["amount"] < 0))
    total_credit = _round_money(sum(t["amount"] for t in filtered if t["amount"] > 0))

    headers = ["Date", "Category", "Type", "Payee", "Debit", "Credit", "Balance"]
    body = _statement_rows(filtered, balances)

    wb = Workbook()

    _balance_sheet_sheet(
        wb.active,
        account,
        search,
        opening_balance,
        closing_balance,
        total_debit,
        total_credit,
        len(body),
        [t for t in filtered if t["category_type"] in ("asset", "liability") or t["has_split"]],
    )

    _transactions_sheet(
        wb.create_sheet("Cashflow"),
        headers,
        _statement_rows(
            [t for t in filtered if t["category_type"] in ("income", "expense")],
            balances,
        ),
    )

    _category_payee_sheet(wb.create_sheet("By Category"), filtered)

    _transactions_sheet(wb.create_sheet("All"), headers, body)

    _style(wb)
    _style_account_statement_columns(wb)

    return _save_workbook(wb, stem)


def _style_account_statement_columns(wb):
    for sheet_name in ("All", "Cashflow"):
        if sheet_name not in wb.sheetnames:
            continue

        sheet = wb[sheet_name]

        payee_width = max(
            (len(str(sheet.cell(row=row, column=4).value or "")) for row in range(1, sheet.max_row + 1)),
            default=12,
        )

        sheet.column_dimensions["A"].width = 14  # Date
        sheet.column_dimensions["B"].width = 28  # Category
        sheet.column_dimensions["C"].width = 12  # Type
        sheet.column_dimensions["D"].width = min(max(payee_width + 2, 12), 30)  # Payee
        sheet.column_dimensions["E"].width = 14  # Debit
        sheet.column_dimensions["F"].width = 14  # Credit
        sheet.column_dimensions["G"].width = 16  # Balance


def _statement_rows(transactions, balances):
    return [
        (
            [
                t["date"],
                t["category"],
                t["category_type"].capitalize() if t["category_type"] else "",
                t["payee"],
                t["amount"] if t["amount"] < 0 else "",
                t["amount"] if t["amount"] > 0 else "",
                balances[t["transaction_hash"]],
            ],
            bool(t["has_split"]),
        )
        for t in transactions
    ]


def _expand_splits(transactions):
    for t in transactions:
        if not t["has_split"]:
            yield t
            continue

        for s in t["splits"]:
            yield {
                "date": t["date"],
                "transaction_hash": t["transaction_hash"],
                "category": s["account"],
                "category_type": s["type"],
                "payee": t["payee"],
                "amount": s["amount"],
                "has_split": False,
            }


def _category_payee_sheet(sheet, transactions):
    sheet.title = "By Category"

    groups = defaultdict(lambda: defaultdict(lambda: [0.0, 0]))
    for t in _expand_splits(transactions):
        category = t["category"] or "Uncategorized"
        payee = t["payee"] or "Unknown"
        entry = groups[category][payee]
        entry[0] += t["amount"]
        entry[1] += 1

    sheet.append(["Category / Payee", "Amount", "Transactions"])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet["B1"].alignment = Alignment(horizontal="right")
    sheet["C1"].alignment = Alignment(horizontal="right")

    grand_total = 0.0
    grand_count = 0

    for category in sorted(groups):
        payees = groups[category]
        total_amount = 0.0
        total_count = 0

        sheet.append([category, "", ""])
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = CATEGORY_FILL

        for payee in sorted(payees):
            amount, count = payees[payee]
            sheet.append([payee, _round_money(amount), count])
            row = sheet.max_row
            sheet.cell(row=row, column=1).alignment = Alignment(indent=1)
            sheet.cell(row=row, column=2).number_format = ACCOUNTING_FORMAT
            sheet.cell(row=row, column=2).alignment = Alignment(horizontal="right")
            sheet.cell(row=row, column=3).alignment = Alignment(horizontal="right")
            total_amount += amount
            total_count += count

        sheet.append(["Total", _round_money(total_amount), total_count])
        total_row = sheet.max_row
        for cell in sheet[total_row]:
            cell.font = Font(bold=True, italic=True)
            cell.border = SUBTOTAL_BORDER
        sheet.cell(row=total_row, column=1).alignment = Alignment(indent=1)
        sheet.cell(row=total_row, column=2).number_format = ACCOUNTING_FORMAT
        sheet.cell(row=total_row, column=2).alignment = Alignment(horizontal="right")
        sheet.cell(row=total_row, column=3).alignment = Alignment(horizontal="right")

        grand_total += total_amount
        grand_count += total_count

        sheet.append([])

    sheet.append(["Grand Total", _round_money(grand_total), grand_count])
    grand_row = sheet.max_row
    for cell in sheet[grand_row]:
        cell.font = Font(bold=True, size=12)
        cell.border = GRAND_TOTAL_BORDER
    sheet.cell(row=grand_row, column=2).number_format = ACCOUNTING_FORMAT
    sheet.cell(row=grand_row, column=2).alignment = Alignment(horizontal="right")
    sheet.cell(row=grand_row, column=3).alignment = Alignment(horizontal="right")

    max_category_len = max((len(c) for c in groups), default=0)
    max_payee_len = max(
        (len(m) for payees in groups.values() for m in payees),
        default=0,
    )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:C{grand_row}"
    sheet.column_dimensions["A"].width = min(max(max(max_category_len, max_payee_len) + 3, 12), 32)
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 14


def _balance_sheet_sheet(
    sheet, account, search, opening_balance, closing_balance, total_debit, total_credit, count, transactions
):
    sheet.title = "Balance Sheet"
    sheet.sheet_properties.outlinePr.summaryBelow = True

    info_pairs = [
        ("Account", account["name"], "Institution", account["institution"] or ""),
        (
            "Account Number",
            f"XXXX{account['last4']}" if account["last4"] else "",
            "Type",
            account["type"].capitalize(),
        ),
        ("Filter", search or "All transactions", "Generated", today().isoformat()),
        ("Opening Balance", opening_balance, "Closing Balance", closing_balance),
        ("Total Debits", total_debit, "Total Credits", total_credit),
        ("Net Change", _round_money(closing_balance - opening_balance), "Transactions", count),
    ]
    for label1, value1, label2, value2 in info_pairs:
        sheet.append([label1, value1, "", label2, value2])
        row = sheet.max_row
        sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.cell(row=row, column=4).font = Font(bold=True)
    sheet.append([])

    assets = defaultdict(list)
    liabilities = defaultdict(list)

    for t in _expand_splits(transactions):
        if t["category_type"] not in ("asset", "liability"):
            continue
        bucket = assets if t["category_type"] == "asset" else liabilities
        bucket[t["category"] or "Unknown"].append(t)

    header_row = len(info_pairs) + 2
    sheet.append(["Date", "Debit", "Credit"])
    for cell in sheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    def _section(title, accounts):
        sheet.append([title, "", ""])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=12)

        section_total = 0.0

        for account_name in sorted(accounts):
            entries = sorted(accounts[account_name], key=lambda t: t["date"])

            sheet.append([account_name, "", ""])
            header_row = sheet.max_row
            sheet.merge_cells(start_row=header_row, start_column=1, end_row=header_row, end_column=3)
            for cell in sheet[header_row]:
                cell.font = Font(bold=True)
                cell.fill = CATEGORY_FILL

            account_total = 0.0
            for t in entries:
                amount = -t["amount"]
                row_values = [
                    t["date"],
                    amount if amount < 0 else "",
                    amount if amount > 0 else "",
                ]
                sheet.append(row_values)
                row = sheet.max_row
                sheet.row_dimensions[row].outlineLevel = 1
                sheet.cell(row=row, column=2).number_format = ACCOUNTING_FORMAT
                sheet.cell(row=row, column=3).number_format = ACCOUNTING_FORMAT
                account_total += amount

            sheet.append(["Subtotal", "", _round_money(account_total)])
            sub_row = sheet.max_row
            for cell in sheet[sub_row]:
                cell.font = Font(bold=True, italic=True)
                cell.border = SUBTOTAL_BORDER
            sheet.cell(row=sub_row, column=3).number_format = ACCOUNTING_FORMAT

            section_total += account_total
            sheet.append([])

        sheet.append([f"Total {title}", "", _round_money(section_total)])
        total_row = sheet.max_row
        for cell in sheet[total_row]:
            cell.font = Font(bold=True)
            cell.border = SUBTOTAL_BORDER
        sheet.cell(row=total_row, column=3).number_format = ACCOUNTING_FORMAT
        sheet.append([])

        return section_total

    assets_total = _section("Assets", assets)
    liabilities_total = _section("Liabilities", liabilities)

    sheet.append(["Net", "", _round_money(assets_total - liabilities_total)])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True, size=12)
        cell.border = GRAND_TOTAL_BORDER
    sheet.cell(row=sheet.max_row, column=3).number_format = ACCOUNTING_FORMAT

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 16


def export_project_report(project_id):
    project = get_project(project_id)
    if project is None:
        raise ValueError(f"Project {project_id} not found.")

    stem = f"project-report_{_slugify(project['name'])}"

    summary = get_project_summary(project_id)
    categories = _positive_amounts(get_project_category_summary(project_id))
    payees = _positive_amounts(get_project_payee_summary(project_id))
    inflows = get_project_inflow_summary(project_id)
    transactions = get_project_transactions(project_id)

    wb = Workbook()

    _project_summary_sheet(wb.active, project, summary, categories)

    txn_sheet = wb.create_sheet("Transactions")
    _transactions_sheet(txn_sheet, TRANSACTIONS_HEADERS, transactions)
    _collapse_column(txn_sheet, len(TRANSACTIONS_HEADERS))

    _table_sheet(
        wb.create_sheet("Payees"),
        ["Payee", "Amount", "Transactions", "%"],
        [
            [
                x["payee"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in payees
        ],
    )

    _table_sheet(
        wb.create_sheet("Inflows"),
        ["Source", "Amount", "Transactions", "%"],
        [
            [
                x["source"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in inflows
        ],
    )

    _style(wb)

    return _save_workbook(wb, stem)


def _project_summary_sheet(sheet, project, summary, categories):
    sheet.title = "Summary"

    budget_amount = Money(project["budget_cents"]).amount if project["budget_cents"] else None

    rows = [
        ["Project", project["name"]],
        ["Type", project["type"] or ""],
        ["Status", "Archived" if project["archived"] else "Active"],
        ["Start Date", project["start_date"] or ""],
        ["End Date", project["end_date"] or "Ongoing"],
        ["Notes", project["notes"] or ""],
        ["Generated", today().isoformat()],
        [],
        ["Transactions", summary["transactions"]],
        ["Income", summary["income"]],
        ["Expenses", summary["expenses"]],
        ["Net", summary["net"]],
    ]

    if budget_amount:
        rows.append(["Budget", budget_amount])
        rows.append(
            [
                "Budget Used %",
                round(100 * summary["expenses"] / budget_amount, 1),
            ]
        )

    rows.append([])
    rows.append(["Category", "Amount", "Transactions", "%"])
    rows.extend(
        [
            [
                x["category"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in categories
        ]
    )

    for row in rows:
        sheet.append(row)


def _slugify(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def _summary_sheet(sheet, period, summary, categories):
    sheet.title = "Summary"

    rows = [
        ["Period", period],
        ["Generated", today().isoformat()],
        [],
        ["Transactions", summary["transactions"]],
        ["Income", summary["income"]],
        ["Expenses", summary["expenses"]],
        ["Net", summary["net"]],
        ["Cash In", summary["inflow"]],
        ["Cash Out", summary["outflow"]],
        ["Cash Flow", summary["cashflow"]],
        [],
        ["Category", "Amount", "Transactions", "%"],
    ]

    rows.extend(
        [
            [
                x["category"],
                x["amount"],
                x["transactions"],
                x["percent"],
            ]
            for x in categories
        ]
    )

    for row in rows:
        sheet.append(row)


def _table_sheet(sheet, headers, rows):
    sheet.append(headers)

    for row in rows:
        sheet.append(row)


def _transactions_sheet(sheet, headers, rows: Iterable[tuple[list, bool]]):
    sheet.append(headers)

    for values, is_split in rows:
        sheet.append(values)

        if is_split:
            for cell in sheet[sheet.max_row]:
                cell.fill = SPLIT_FILL


def _transactions(start, end):
    with db.transaction() as conn:
        rows = conn.execute(
            """
            SELECT
                t.posted_date date,
                COALESCE(
                    m.canonical_name,
                    'Unknown'
                ) payee,
                a.name category,
                a.type account_type,
                ROUND(
                    -l.amount_cents / 100.0,
                    2
                ) amount,
                t.raw_description description,
                EXISTS (
                    SELECT 1 FROM transactions_splits ts
                    WHERE ts.transaction_hash = t.transaction_hash
                ) is_split
            FROM ledger l
            JOIN transactions t
                ON t.id = l.transaction_id
            JOIN accounts a
                ON a.id = l.account_id
            LEFT JOIN payees m
                ON m.id = t.payee_id
            WHERE t.posted_date BETWEEN ? AND ?
              AND l.account_id != t.account_id
              AND a.type IN ('income', 'expense')
            ORDER BY t.posted_date, t.id, l.id
            """,
            [start, end],
        ).fetchall()

    return [
        (
            [
                r["date"],
                r["payee"],
                r["category"],
                r["account_type"].capitalize(),
                r["amount"] if r["amount"] < 0 else "",
                r["amount"] if r["amount"] > 0 else "",
                r["description"],
            ],
            bool(r["is_split"]),
        )
        for r in rows
    ]


def _style(wb):
    for sheet in wb.worksheets:
        is_report = sheet.title in ("By Category", "Balance Sheet")

        if not is_report:
            for cell in sheet[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float) and cell.number_format == "General":
                    cell.number_format = "#,##0.00"

        if is_report:
            continue

        for column in sheet.columns:
            if column[0].value == "%":
                for cell in column[1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '0.0"%"'

            width = max(len(str(cell.value or "")) for cell in column)

            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                max(width + 3, 12),
                80,
            )

    if "All" in wb.sheetnames and "Cashflow" in wb.sheetnames:
        all_widths = wb["All"].column_dimensions
        cashflow_widths = wb["Cashflow"].column_dimensions
        for letter, dim in all_widths.items():
            cashflow_widths[letter].width = dim.width
