import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import xlrd
from openpyxl import load_workbook

from intent_ledger.importer.parsers.base import ParsedStatement, ParsedTransaction

REQUIRED_COLUMNS = ("date", "description", "withdrawal", "deposit")
OPTIONAL_COLUMNS = ("balance",)
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


class CanonicalParser:
    slug = "canonical"
    display_name = "Canonical CSV/XLS/XLSX template"
    file_extensions = (".csv", ".xls", ".xlsx")

    def sniff(self, path: Path) -> bool:
        try:
            rows = _read_rows(path)

            if not rows:
                return False

            header = {_normalize_header(cell) for cell in rows[0]}
            missing = set(REQUIRED_COLUMNS) - header

            return not missing

        except Exception:
            return False

    def parse(self, path: Path) -> ParsedStatement:
        rows = _read_rows(path)

        if not rows:
            raise ValueError("Statement is empty.")

        header = [_normalize_header(cell) for cell in rows[0]]
        missing = set(REQUIRED_COLUMNS) - set(header)

        if missing:
            raise ValueError(f"Missing required column(s): {', '.join(sorted(missing))}")

        column_index = {name: header.index(name) for name in ALL_COLUMNS if name in header}

        transactions: list[ParsedTransaction] = []
        row_errors: list[str] = []

        for row_number, row in enumerate(rows[1:], start=2):
            if not any(cell.strip() for cell in row if cell is not None):
                continue

            try:
                transactions.append(_parse_row(row, column_index))
            except ValueError as exc:
                row_errors.append(f"row {row_number}: {exc}")

        return {"transactions": transactions, "row_errors": row_errors}


def _normalize_header(cell: str) -> str:
    return str(cell or "").strip().lower()


def _read_rows(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()

    if suffix == ".xlsx":
        return _read_xlsx_rows(path)
    if suffix == ".xls":
        return _read_xls_rows(path)

    return _read_csv_rows(path)


def _read_csv_rows(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.reader(handle))


def _read_xlsx_rows(path: Path) -> list[list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        sheet = workbook.active
        rows = [
            [cell if cell is not None else "" for cell in row] for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()

    return [[str(cell) for cell in row] for row in rows]


def _read_xls_rows(path: Path) -> list[list[str]]:
    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)

    return [
        [_xls_cell_to_str(cell, workbook.datemode) for cell in sheet.row(row_index)]
        for row_index in range(sheet.nrows)
    ]


def _xls_cell_to_str(cell, datemode: int) -> str:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""

    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, datemode).strftime("%Y-%m-%d")

    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return str(int(cell.value)) if cell.value == int(cell.value) else str(cell.value)

    return str(cell.value)


def _parse_row(row: list[str], column_index: dict[str, int]) -> ParsedTransaction:
    posted_date = _parse_date(_cell(row, column_index, "date"))
    raw_description = _cell(row, column_index, "description").strip()
    withdrawal_cents = _parse_decimal_cents(_cell(row, column_index, "withdrawal"), "withdrawal")
    deposit_cents = _parse_decimal_cents(_cell(row, column_index, "deposit"), "deposit")
    balance_cents = _parse_decimal_cents(_cell(row, column_index, "balance"), "balance")

    if not raw_description:
        raise ValueError("description is required.")

    if withdrawal_cents and deposit_cents:
        raise ValueError("row cannot have both a withdrawal and a deposit amount.")
    if not withdrawal_cents and not deposit_cents:
        raise ValueError("row must have a withdrawal or deposit amount.")

    amount_cents = (deposit_cents or 0) - (withdrawal_cents or 0)

    return {
        "posted_date": posted_date,
        "raw_description": raw_description,
        "amount_cents": amount_cents,
        "balance_cents": balance_cents,
    }


def _cell(row: list[str], column_index: dict[str, int], name: str) -> str:
    index = column_index.get(name)

    if index is None or index >= len(row):
        return ""

    return row[index] or ""


def _parse_date(raw: str) -> date:
    raw = str(raw).strip()

    if not raw:
        raise ValueError("date is required.")

    formats = (
        # ISO
        "%Y-%m-%d",
        "%Y/%m/%d",
        # Day-first numeric
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%d.%m.%Y",
        # Day-first textual
        "%d %b %Y",
        "%d %B %Y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        # Month-first numeric
        "%m-%d-%Y",
        "%m/%d/%Y",
        # Other common formats
        "%Y%m%d",
    )

    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue

    raise ValueError(
        f"invalid date '{raw}'. "
        f"Supported formats include YYYY-MM-DD, DD/MM/YYYY, "
        f"DD-MM-YYYY, DD Mon YYYY, and MM/DD/YYYY."
    )


def _parse_decimal_cents(raw: str, field_name: str) -> int | None:
    raw = raw.strip()

    if not raw:
        return None

    try:
        value = Decimal(raw)
    except InvalidOperation:
        raise ValueError(f"invalid {field_name} '{raw}'.") from None

    if -value.as_tuple().exponent > 2:
        raise ValueError(f"{field_name} '{raw}' has more than 2 decimal places.")

    return int(value * 100)
