import csv
import zipfile

import pytest
from openpyxl import load_workbook

from intent_ledger.accounting.ledger import rebuild_ledger
from intent_ledger.accounting.projects import assign_transactions_to_project, create_project
from intent_ledger.accounting.rules_overrides import add_override
from intent_ledger.analytics import workbook
from intent_ledger.analytics.workbook import (
    _expand_splits,
    _round_money,
    _slugify,
    export_account_statement,
    export_monthly_report,
    export_project_report,
    export_yearly_report,
)
from intent_ledger.settings import set_export_format


def csv_rows(path, sheet_name):
    with zipfile.ZipFile(path) as archive:
        return list(csv.reader(archive.read(f"{sheet_name}.csv").decode().splitlines()))


class FakeForm(dict):
    def get(self, key, default=None, type=None):
        if key not in self:
            return default
        value = super().get(key)
        if type is None:
            return value
        try:
            return type(value)
        except (TypeError, ValueError):
            return default


@pytest.fixture(autouse=True)
def redirect_exports(tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "FINANCE_EXPORT_DIR", tmp_path / "exports")


def sheet_rows(path, sheet_name):
    wb = load_workbook(path)
    return [tuple(row) for row in wb[sheet_name].iter_rows(values_only=True)]


def test_round_money_normalizes_negative_zero_to_positive_zero():
    assert _round_money(-0.001) == 0.0
    assert str(_round_money(-0.001)) == "0.0"
    assert _round_money(12.345) == 12.35
    assert _round_money(-12.345) == -12.35


def test_slugify_lowercases_and_replaces_non_alphanumerics():
    assert _slugify("Checking Account #1") == "checking-account-1"
    assert _slugify("  Leading/Trailing  ") == "leading-trailing"


def test_expand_splits_passes_through_non_split_and_expands_split_rows():
    rows = [
        {"has_split": False, "date": "2026-01-01", "category": "Groceries", "amount": -10.0},
        {
            "has_split": True,
            "date": "2026-01-02",
            "transaction_hash": "h1",
            "payee": "Costco",
            "splits": [
                {"account": "Groceries", "type": "expense", "amount": -6.0},
                {"account": "Household", "type": "expense", "amount": -4.0},
            ],
        },
    ]

    expanded = list(_expand_splits(rows))

    assert len(expanded) == 3
    assert expanded[0]["category"] == "Groceries" and expanded[0]["amount"] == -10.0
    assert expanded[1] == {
        "date": "2026-01-02",
        "transaction_hash": "h1",
        "category": "Groceries",
        "category_type": "expense",
        "payee": "Costco",
        "amount": -6.0,
        "has_split": False,
    }
    assert expanded[2]["category"] == "Household" and expanded[2]["amount"] == -4.0


def test_export_account_statement_running_balance_uses_full_unfiltered_ledger(
    conn, account_factory, transaction_factory
):
    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'GROCERY RUN', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    transaction_factory(checking_id, "2026-01-01", 100000, "OPENING DEPOSIT")
    transaction_factory(checking_id, "2026-01-05", -1500, "GROCERY RUN")
    transaction_factory(checking_id, "2026-01-10", -2500, "GROCERY RUN")
    rebuild_ledger()

    path = export_account_statement(checking_id, search="category:groceries")

    rows = sheet_rows(path, "All")
    header, *data_rows = rows

    assert header == ("Date", "Category", "Type", "Payee", "Debit", "Credit", "Balance")
    assert data_rows[0][0] == "2026-01-05"
    assert data_rows[0][4] == -15.0
    assert data_rows[0][6] == pytest.approx(985.0)
    assert data_rows[1][0] == "2026-01-10"
    assert data_rows[1][6] == pytest.approx(960.0)


def test_export_account_statement_excludes_equity_from_filtered_rows_but_not_balance(
    conn, account_factory, transaction_factory
):
    checking_id = account_factory("Test Checking")
    equity_id = account_factory("Test Opening Balances", type="equity")

    opening_hash = transaction_factory(checking_id, "2026-01-01", 50000, "Open")
    transaction_factory(checking_id, "2026-01-05", -2000, "COFFEE")
    rebuild_ledger()

    add_override(FakeForm({"transaction_hash": opening_hash, "account_id": equity_id}))
    rebuild_ledger()

    path = export_account_statement(checking_id)

    rows = sheet_rows(path, "All")
    _header, *data_rows = rows

    assert len(data_rows) == 1
    assert data_rows[0][0] == "2026-01-05"
    assert data_rows[0][6] == pytest.approx(480.0)

    summary_rows = sheet_rows(path, "Balance Sheet")
    summary = {}
    for r in summary_rows:
        if r and r[0]:
            summary[r[0]] = r[1]
        if len(r) > 3 and r[3]:
            summary[r[3]] = r[4]
    assert summary["Opening Balance"] == pytest.approx(500.0)
    assert summary["Closing Balance"] == pytest.approx(480.0)


def test_export_account_statement_raises_for_unknown_account(conn):
    with pytest.raises(ValueError):
        export_account_statement(999999)


def test_export_account_statement_by_category_sheet_totals_match_transactions(
    conn, account_factory, transaction_factory
):
    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'GROCERY RUN', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    transaction_factory(checking_id, "2026-01-05", -1500, "GROCERY RUN")
    transaction_factory(checking_id, "2026-01-10", -2500, "GROCERY RUN")
    rebuild_ledger()

    path = export_account_statement(checking_id)

    rows = sheet_rows(path, "By Category")
    row_by_first_cell = {r[0]: r for r in rows if r[0]}
    assert row_by_first_cell["Grand Total"][1] == pytest.approx(-40.0)


def test_export_monthly_report_has_expected_sheets_and_category_totals(
    conn, account_factory, transaction_factory
):
    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'GROCERY RUN', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    transaction_factory(checking_id, "2026-03-05", -1500, "GROCERY RUN")
    rebuild_ledger()

    path = export_monthly_report(2026, 3)

    wb = load_workbook(path)
    assert set(wb.sheetnames) == {"Summary", "Transactions", "Payees", "Inflows"}

    summary = sheet_rows(path, "Summary")
    assert ("Test Groceries", 15.0, 1, 100.0) in summary


def test_export_monthly_report_as_csv_zips_one_csv_per_sheet(conn, account_factory, transaction_factory):
    set_export_format(conn, "csv")
    conn.commit()

    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'GROCERY RUN', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    transaction_factory(checking_id, "2026-03-05", -1500, "GROCERY RUN")
    rebuild_ledger()

    path = export_monthly_report(2026, 3)

    assert path.suffix == ".zip"
    with zipfile.ZipFile(path) as archive:
        assert set(archive.namelist()) == {"summary.csv", "transactions.csv", "payees.csv", "inflows.csv"}

    assert ["Test Groceries", "15.0", "1", "100.0"] in csv_rows(path, "summary")


def test_export_account_statement_as_csv_preserves_all_sheets(conn, account_factory, transaction_factory):
    set_export_format(conn, "csv")
    conn.commit()

    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'GROCERY RUN', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    transaction_factory(checking_id, "2026-01-05", -1500, "GROCERY RUN")
    rebuild_ledger()

    path = export_account_statement(checking_id)

    with zipfile.ZipFile(path) as archive:
        assert set(archive.namelist()) == {"balance-sheet.csv", "cashflow.csv", "by-category.csv", "all.csv"}

    header, *data_rows = csv_rows(path, "all")
    assert header == ["Date", "Category", "Type", "Payee", "Debit", "Credit", "Balance"]
    assert data_rows[0][0] == "2026-01-05"


def test_export_yearly_report_covers_the_full_year(conn, account_factory, transaction_factory):
    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'GROCERY RUN', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    transaction_factory(checking_id, "2026-01-05", -1500, "GROCERY RUN")
    transaction_factory(checking_id, "2026-11-05", -2500, "GROCERY RUN")
    rebuild_ledger()

    path = export_yearly_report(2026)

    txns = sheet_rows(path, "Transactions")
    dates = {row[0] for row in txns[1:]}
    assert dates == {"2026-01-05", "2026-11-05"}


def test_export_project_report_includes_only_assigned_transactions(
    conn, account_factory, transaction_factory
):
    checking_id = account_factory("Test Checking")
    groceries_id = account_factory("Test Groceries", type="expense")
    conn.execute(
        "INSERT INTO account_rules (match_type, pattern, account_id, priority) "
        "VALUES ('equals', 'PROJECT SPEND', ?, 0)",
        (groceries_id,),
    )
    conn.commit()

    in_project_hash = transaction_factory(checking_id, "2026-01-05", -5000, "PROJECT SPEND")
    transaction_factory(checking_id, "2026-01-06", -3000, "PROJECT SPEND")
    rebuild_ledger()

    project_id = create_project(FakeForm({"name": "Kitchen Remodel"}))
    assign_transactions_to_project([in_project_hash], project_id)

    path = export_project_report(project_id)

    txns = sheet_rows(path, "Transactions")
    assert len(txns) == 2
    assert txns[1][0] == "2026-01-05"


def test_export_project_report_raises_for_unknown_project(conn):
    with pytest.raises(ValueError):
        export_project_report(999999)
